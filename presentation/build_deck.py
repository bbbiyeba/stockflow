"""Build the Stryker presentation (PPTX) from presentation/assets (run make_charts.py first)."""

from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
A = HERE / "assets"
F = json.loads((A / "facts.json").read_text())

GOLD, BLACK, CHARCOAL, GRAY = (
    RGBColor(0xFF, 0xB5, 0x02),
    RGBColor(0, 0, 0),
    RGBColor(0x33, 0x33, 0x33),
    RGBColor(0x6E, 0x6E, 0x6E),
)
LIGHT, LIGHTER, WHITE = RGBColor(0xD9, 0xD9, 0xD9), RGBColor(0xF2, 0xF2, 0xF2), RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)
ML = Inches(0.6)  # left/right margin
CW = W - 2 * ML  # content width
BODY_TOP = Inches(1.75)
FOOT_Y = Inches(7.0)
DEMO_URL = "stockflow-six-mu.vercel.app"


def money(x: float, d: int = 1) -> str:
    """Format dollars as $1.2M, or $45k below a million."""
    return f"${x / 1e6:.{d}f}M" if abs(x) >= 1e6 else f"${x / 1e3:,.0f}k"


def pct(x: float, d: int = 0) -> str:
    """Format a fraction as a percentage with ``d`` decimals."""
    return f"{x * 100:.{d}f}%"


# --------------------------------------------------------------------------- primitives
def text(
    slide,
    x,
    y,
    w,
    h,
    runs,
    size=14,
    color=CHARCOAL,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing=1.1,
):
    """Text box. ``runs`` is a str or a list of paragraphs, each a str or list of (text, overrides)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        for chunk in (para if isinstance(para, list) else [(para, {})]):
            t, o = (chunk, {}) if isinstance(chunk, str) else chunk
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def _flat(shape):
    """Drop the theme style reference (it adds a shadow in some renderers); we set fill/line explicitly."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    """Flat filled rectangle (no shadow)."""
    s = _flat(slide.shapes.add_shape(shape, x, y, w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def box(slide, x, y, w, h, label, fill=WHITE, line=CHARCOAL, size=13, bold=True, color=BLACK):
    """Outlined box with centred, possibly multi-line, label."""
    s = rect(slide, x, y, w, h, fill, line)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line_text in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line_text
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return s


def arrow(slide, x1, y1, x2, y2, color=CHARCOAL, width=1.5):
    """Flat straight connector with an arrow head."""
    c = _flat(slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def picture(slide, path, x, y, w=None, h=None):
    """Place an image fitted inside (w, h), centred in that box, preserving aspect ratio."""
    iw, ih = Image.open(path).size
    if w is not None and h is not None:
        scale = min(w / iw, h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        return slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)
    return slide.shapes.add_picture(str(path), x, y, w, h)


def table(slide, x, y, w, rows, col_w, header=True, size=12, row_h=Inches(0.42), align=None, bold_last=False):
    """Plain table: grey header, light rules, explicit alignment."""
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, x, y, w, row_h * nr)
    t = gt.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):  # plain styling, our own colours
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for j, cw in enumerate(col_w):
        t.columns[j].width = cw
    for i, row in enumerate(rows):
        t.rows[i].height = row_h
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHTER if (header and i == 0) else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = align[j] if align else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(val)
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold = (header and i == 0) or (bold_last and i == nr - 1)
            r.font.color.rgb = CHARCOAL if (header and i == 0) else BLACK
            # thin light-gray horizontal rules only
            tcPr = cell._tc.get_or_add_tcPr()
            for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcPr.find(qn(edge))
                if ln is not None:
                    tcPr.remove(ln)
            for edge, on in (("a:lnL", False), ("a:lnR", False), ("a:lnT", i == 0), ("a:lnB", True)):
                ln = tcPr.makeelement(qn(edge), {"w": str(Pt(1 if on else 0)), "cmpd": "sng"})
                if on:
                    sf = ln.makeelement(qn("a:solidFill"), {})
                    clr = sf.makeelement(qn("a:srgbClr"), {"val": "D9D9D9"})
                    sf.append(clr)
                    ln.append(sf)
                else:
                    ln.append(ln.makeelement(qn("a:noFill"), {}))
                tcPr.append(ln)
    return gt


def callout(slide, n, cx, cy, d=Inches(0.42)):
    """Numbered gold marker (black numeral) centred on (cx, cy)."""
    s = rect(slide, cx - d // 2, cy - d // 2, d, d, GOLD, BLACK, MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(14), True, BLACK


def frame(prs, title, n, notes):
    """New slide with accent bar, takeaway title, footer and speaker notes."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, ML, Inches(0.42), Inches(0.9), Inches(0.09), GOLD)
    text(
        s,
        ML,
        Inches(0.62),
        CW,
        Inches(1.0),
        title,
        size=26,
        bold=True,
        color=BLACK,
        anchor=MSO_ANCHOR.TOP,
        line_spacing=1.0,
    )
    rect(s, ML, FOOT_Y - Inches(0.12), CW, Emu(9525), LIGHT)
    text(s, ML, FOOT_Y, Inches(6), Inches(0.3), "Synthetic data – for demonstration", size=10, color=GRAY)
    text(s, W - ML - Inches(1), FOOT_Y, Inches(1), Inches(0.3), str(n), size=10, color=GRAY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes.strip()
    return s


# --------------------------------------------------------------------------- slides
def s1_title(prs):
    """Slide 1: title, author and mission line."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, ML, Inches(1.55), Inches(0.9), Inches(0.09), GOLD)
    text(s, ML, Inches(1.85), Inches(6.4), Inches(1.6), "Stockflow", size=44, bold=True, color=BLACK)
    text(
        s,
        ML,
        Inches(2.75),
        Inches(6.2),
        Inches(1.0),
        "Branch-level demand forecasting and inventory rebalancing",
        size=22,
        color=CHARCOAL,
    )
    text(
        s,
        ML,
        Inches(4.0),
        Inches(6.0),
        Inches(1.0),
        "Right product at the right branch, so surgeons, hospitals and their patients are not left waiting.",
        size=16,
        color=CHARCOAL,
    )
    rect(s, ML, Inches(5.25), Inches(5.6), Emu(9525), LIGHT)
    text(
        s,
        ML,
        Inches(5.45),
        Inches(6),
        Inches(1.0),
        [[("Bryce Biyeba", {"bold": True, "color": BLACK})], "Data Science Intern  ·  October 2026"],
        size=15,
    )
    picture(s, A / "branch_map.png", Inches(7.0), Inches(1.4), Inches(5.8), Inches(4.4))
    text(
        s,
        Inches(7.0),
        Inches(5.85),
        Inches(5.8),
        Inches(0.4),
        f"{F['branches']} branches · {F['skus']} SKUs · {F['series']:,} branch × SKU combinations",
        size=12,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )
    rect(s, ML, FOOT_Y - Inches(0.12), CW, Emu(9525), LIGHT)
    text(s, ML, FOOT_Y, Inches(6), Inches(0.3), "Synthetic data – for demonstration", size=10, color=GRAY)
    text(s, W - ML - Inches(1), FOOT_Y, Inches(1), Inches(0.3), "1", size=10, color=GRAY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = f"""
Good morning. I'm Bryce Biyeba, a data science intern, and this is Stockflow: a tool that forecasts demand for every product at every branch and recommends moving stock between branches before anyone runs out.

Why it matters for Stryker: when an implant or a piece of equipment isn't at the branch that needs it, a surgeon waits, and so does the patient. Our mission is to make healthcare better together with our customers, and having the right product at the right branch is a very concrete part of that.

One thing to say up front, in the spirit of integrity: everything you'll see runs on synthetic data I generated to behave like a real distributor's ERP and warehouse extracts: {F['branches']} branches, {F['skus']} SKUs, three years of daily history, deliberately messy. The method is real; the dollar figures describe that simulated network, not Stryker.

In the next eight minutes: the problem, how the system works, how accurate the forecasts are, what the rebalancing engine recommends, the business impact, what the dashboard looks like for a branch manager, and a proposal for a pilot.

Likely question: "Why synthetic data?" Answer: I didn't have access to company data, so I built a generator that includes the problems real extracts have (outages, keying errors, stock-outs) so the cleaning and modeling had to handle them for real.
"""


def s2_problem(prs):
    """Slide 2: the inventory problem in three numbers and a chart."""
    h = F["health"]
    s = frame(
        prs,
        f"{pct(F['excess_share'])} of inventory dollars sit above target while "
        f"{F['risk_positions']} branch positions face stock-outs",
        2,
        f"""
Here's the problem in one picture. Across the network there's {money(F['on_hand_value'])} of inventory on hand. {money(F['excess_value'])} of it, {pct(F['excess_share'])}, sits above the level each branch actually needs. At the same time, {F['risk_positions']} branch-and-product combinations are likely to run out before their next replenishment arrives.

Over three years of history, {pct(F['stockout_day_rate'], 1)} of all branch-product days were stock-out days, days when a customer could have asked for something the branch didn't have.

The chart shows where the money is. The gold bar is the inventory at the positions facing stock-outs: only {money(h['stockout_risk']['value'])}. The shortage isn't a lack of inventory; it's inventory in the wrong place. A branch in one city has months of supply while another is about to stock out of the same item.

And {money(F['stagnant_value'])} hasn't sold in more than 90 days. Almost all of that, {money(F['liquidation_value'])}, is product the whole network can't sell within a year, so that's a liquidation decision, not something transfers can fix. I call that out separately so we don't pretend the model solves it.

Likely question: "How is 'above target' defined?" Answer: target is forecast demand over the lead time plus the review period, plus safety stock at a 95% service level. Anything above that is excess.
""",
    )
    stats = [
        (pct(F["stockout_day_rate"], 1), "of branch × SKU days were stock-out days (3 years)"),
        (money(F["excess_value"]), f"above target, of {money(F['on_hand_value'])} on hand"),
        (money(F["stagnant_value"]), "stagnant: no sale in 90+ days"),
    ]
    y = BODY_TOP + Inches(0.1)
    for big, small in stats:
        rect(s, ML, y + Inches(0.08), Inches(0.08), Inches(0.95), GOLD)
        text(s, ML + Inches(0.25), y, Inches(4.2), Inches(0.7), big, size=34, bold=True, color=BLACK)
        text(s, ML + Inches(0.25), y + Inches(0.66), Inches(4.2), Inches(0.5), small, size=13, color=CHARCOAL)
        y += Inches(1.6)
    picture(s, A / "problem.png", Inches(5.4), BODY_TOP, Inches(7.33), Inches(4.6))
    text(
        s,
        Inches(5.4),
        Inches(6.45),
        Inches(7.33),
        Inches(0.35),
        "Inventory on hand by position status, at cost",
        size=11,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )


def s3_approach(prs):
    """Slide 3: pipeline architecture."""
    s = frame(
        prs,
        "One pipeline turns three years of branch data into transfers a planner approves",
        3,
        f"""
This is the whole system on one line, left to right.

It starts from the extracts any ERP and warehouse system already produces: daily sales, stock movements, open purchase orders, promotions and lane costs. The first step cleans and validates them: duplicates, returns, system outages, keying errors, and, importantly, days when a branch was out of stock, because sales on those days understate real demand.

Then the forecast: a single machine-learning model, LightGBM, trained across all {F['series']:,} branch-product combinations, alongside classic statistical models for slow movers. The best one is picked per segment. Every forecast comes with a range, not just a number.

The inventory policy turns that range into safety stock, reorder points and a stock-out probability for every position.

The optimizer then finds the cheapest set of transfers that covers shortfalls from branches with excess, respecting safety stock, storage capacity and transit time.

Finally, the dashboard: and this last box in gold matters most. Nothing moves until a planner approves it. Every approval or rejection is logged with who, when and why. That's accountability built in.

Likely question: "How long does a run take?" Answer: about five minutes for this network; I also tested a network ten times larger, which completes on a single 4-core machine.
""",
    )
    steps = [
        ("ERP / WMS\nextracts", "sales, stock moves,\nPOs, promotions"),
        ("Clean &\nvalidate", "outages, returns,\nstock-out days"),
        ("Forecast", "LightGBM + statistical,\nP10 / P50 / P90"),
        ("Inventory\npolicy", "safety stock,\nreorder point, risk"),
        ("Rebalancing\noptimizer", "min-cost transfers\nunder constraints"),
        ("Dashboard", "plain-language\nreason per move"),
        ("Planner\napproval", "nothing moves\nwithout sign-off"),
    ]
    n = len(steps)
    gap = Inches(0.32)
    bw = (CW - gap * (n - 1)) // n
    bh = Inches(1.25)
    y = Inches(2.55)
    for i, (lab, sub) in enumerate(steps):
        x = ML + i * (bw + gap)
        last = i == n - 1
        box(s, x, y, bw, bh, lab, fill=GOLD if last else WHITE, line=BLACK if last else CHARCOAL, size=14)
        text(
            s,
            x - Inches(0.05),
            y + bh + Inches(0.18),
            bw + Inches(0.1),
            Inches(0.8),
            sub.split("\n"),
            size=11,
            color=GRAY,
            align=PP_ALIGN.CENTER,
        )
        if i < n - 1:
            arrow(s, x + bw + Inches(0.03), y + bh // 2, x + bw + gap - Inches(0.03), y + bh // 2)
    text(
        s,
        ML,
        Inches(5.2),
        CW,
        Inches(0.5),
        "Every step checks its inputs; every assumption is documented in the repository.",
        size=13,
        color=CHARCOAL,
        align=PP_ALIGN.CENTER,
    )


def s4_forecast(prs):
    """Slide 4: forecast vs actual and the model comparison."""
    m = F["models"]
    fx = F["forecast_example"]
    improve = 1 - m["selected"]["wape"] / m["seasonal_naive"]["wape"]
    s = frame(
        prs,
        f"Forecasts are {pct(improve)} more accurate than the seasonal baseline on a month the " f"model never saw",
        4,
        f"""
To keep the evaluation honest, I chose the models using three earlier months and then scored them on a fourth month that played no part in any decision. That's the month you see here.

The chart is one real branch-product pair, {fx['branch_id']} and {fx['sku_id']}, a {fx['category']} item. I picked it by rule, not by eye: it's the A-class item whose error is closest to the A-class median. Grey is the month before, black is what actually sold, gold is the forecast, and the shaded band is the 10th-to-90th percentile range. Daily demand is noisy; the forecast tracks the weekly rhythm and the level.

The table compares models on the held-out month. WAPE is total absolute error as a share of actual units. The seasonal-naive baseline, which just repeats last week, misses by {pct(m['seasonal_naive']['wape'])}. The selected models miss by {pct(m['selected']['wape'])}, a {pct(improve)} improvement, with bias under one percent. Over a full 28-day total, which is what drives safety stock, the error is {pct(m['selected']['wape_28d_total'])}.

The P10-to-P90 band contained {pct(F['coverage'])} of actual days against an 80% target: slightly conservative, which I've documented.

Likely question: "Why is daily error still around fifty percent?" Answer: many items sell a few units a day or less, so daily numbers are inherently noisy; that's why we plan on totals and ranges, not single days.
""",
    )
    pic = picture(s, A / "forecast.png", ML, BODY_TOP, Inches(7.0), Inches(4.4))
    text(
        s,
        ML,
        pic.top + pic.height + Inches(0.1),
        Inches(7.0),
        Inches(0.4),
        f"Branch {fx['branch_id']}, SKU {fx['sku_id']} ({fx['category']}), {fx['start']} to {fx['end']}",
        size=11,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )
    rows = [["Model (held-out month)", "WAPE", "MAE", "Bias"]]
    for key, lab in [
        ("seasonal_naive", "Seasonal naive (baseline)"),
        ("ses", "Exponential smoothing"),
        ("croston_sba", "Croston (intermittent)"),
        ("lightgbm", "LightGBM (global ML)"),
        ("selected", "Selected per segment"),
    ]:
        r = m[key]
        rows.append([lab, pct(r["wape"], 1), f"{r['mae']:.2f}", f"{r['bias'] * 100:+.1f}%"])
    R = PP_ALIGN.RIGHT
    table(
        s,
        Inches(7.95),
        BODY_TOP + Inches(0.15),
        Inches(4.78),
        rows,
        [Inches(2.33), Inches(0.75), Inches(1.0), Inches(0.7)],
        size=12,
        row_h=Inches(0.5),
        align=[PP_ALIGN.LEFT, R, R, R],
        bold_last=True,
    )
    text(
        s,
        Inches(7.95),
        Inches(5.35),
        Inches(4.78),
        Inches(0.9),
        [
            f"WAPE = absolute error ÷ actual units; MAE in units per day. P10–P90 band held "
            f"{pct(F['coverage'])} of actual days (target 80%)."
        ],
        size=11,
        color=GRAY,
    )


def s5_rebalance(prs):
    """Slide 5: transfer map and the top five transfers."""
    t = F["top_transfers"]
    s = frame(
        prs,
        f"{F['transfers']} transfers cost {money(F['transfer_cost'])} to move and save "
        f"{money(F['net_savings'])} net, {money(F['net_savings'] - F['greedy_net_savings'])} more than a "
        f"rule of thumb",
        5,
        f"""
This is the rebalancing engine's output. It weighs three costs at once: the freight to move a unit, the lost sale if a branch runs out, and the holding and obsolescence cost of excess left sitting. It also never pulls a sending branch below its own reorder point, never overfills a receiver, respects storage capacity, skips tiny moves, and only uses lanes of four transit days or less.

The map shows the {F['map']['corridors_shown']} busiest corridors, out of {F['map']['corridors_total']}; they carry {pct(F['map']['share_units_shown'])} of the units. Black dots are branches sending stock, gold dots are branches receiving it.

The table is the five most valuable moves. Every one is "excess to shortfall": a branch with more than it needs sending to a branch about to run out. For example, the top move sends {t[0]['units']} units of {t[0]['sku']} from {t[0]['from']} to {t[0]['to']} for {money(t[0]['cost'])} in freight and saves about {money(t[0]['savings'])}.

Overall: {F['transfers']} transfers, {money(F['transfer_cost'])} in freight, {money(F['net_savings'])} net savings, compared with {money(F['greedy_net_savings'])} from a simple rule of thumb that serves the most urgent shortages first.

Likely question: "How stable is the plan?" Answer: I re-solved it with service level, freight cost, forecast error and forecast level each changed. {pct(F['sens_overlap_min'])} to {pct(F['sens_overlap_max'])} of the same moves survive, and net savings stay between {money(F['sens_savings_min'])} and {money(F['sens_savings_max'])}.
""",
    )
    picture(s, A / "transfer_map.png", ML, BODY_TOP, Inches(6.4), Inches(4.3))
    text(
        s,
        ML,
        Inches(6.2),
        Inches(6.4),
        Inches(0.5),
        f"Top {F['map']['corridors_shown']} of {F['map']['corridors_total']} corridors · line width = units · "
        f"black = sending branch, gold = receiving",
        size=11,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )
    rows = [["From → To", "SKU", "Units", "Net savings", "Reason"]]
    for r in t:
        rows.append(
            [
                f"{r['from']} → {r['to']}",
                r["sku"],
                f"{r['units']:,}",
                f"${r['savings']:,.0f}",
                "Excess → shortfall" if r["reason"] == "excess->shortfall" else r["reason"],
            ]
        )
    R, L = PP_ALIGN.RIGHT, PP_ALIGN.LEFT
    table(
        s,
        Inches(7.3),
        BODY_TOP + Inches(0.15),
        Inches(5.43),
        rows,
        [Inches(1.25), Inches(0.8), Inches(0.68), Inches(1.15), Inches(1.55)],
        size=12,
        row_h=Inches(0.5),
        align=[L, L, R, R, L],
    )
    text(
        s,
        Inches(7.3),
        Inches(5.35),
        Inches(5.43),
        Inches(0.9),
        "Each move ships with a plain-language reason a branch manager can check.",
        size=12,
        color=CHARCOAL,
    )


def s6_impact(prs):
    """Slide 6: before / after impact."""
    share = F["stockout_avoided_share"]
    s = frame(
        prs,
        f"Rebalancing frees {money(F['excess_freed'])} of excess and avoids {pct(share)} of expected " f"lost sales",
        6,
        f"""
Here is the impact of executing the recommended plan, compared with doing nothing.

The headline, in gold: {money(F['excess_freed'], 2)} of excess inventory put back to work. Those units go where they'll sell instead of sitting on a shelf, which also means purchases the distribution center doesn't have to make.

For patients this is the line that matters most: expected lost-sale units during the replenishment lead time drop from {F['stockout_units_before']:,.0f} to {F['stockout_units_after']:,.0f}. That's {pct(share)} of the units that would otherwise have been missing when someone asked for them. Positions at stock-out risk fall from {F['risk_positions']} to {F['risk_after']}.

The annual carrying cost of excess falls from {money(F['carry_before'], 2)} to {money(F['carry_after'], 2)}, and inventory turns improve from {F['turns_before']:.1f} to {F['turns_after']:.1f}.

Two honest caveats. First, these are model estimates on synthetic data, so the size of the numbers won't transfer, but the method will. Second, stagnant stock barely changes ({money(F['stagnant_before'])} before, {money(F['stagnant_after'])} after), because nobody else in the network wants it; that needs a liquidation decision.

Likely question: "How are turns measured if transfers don't remove inventory?" Answer: on projected inventory after one replenishment cycle; units moved out of excess replace new purchases.
""",
    )
    panel_w = Inches(4.1)
    rect(s, ML, BODY_TOP + Inches(0.1), panel_w, Inches(4.5), GOLD)
    text(
        s,
        ML + Inches(0.35),
        BODY_TOP + Inches(1.0),
        panel_w - Inches(0.7),
        Inches(1.4),
        money(F["excess_freed"]),
        size=66,
        bold=True,
        color=BLACK,
    )
    text(
        s,
        ML + Inches(0.35),
        BODY_TOP + Inches(2.45),
        panel_w - Inches(0.7),
        Inches(1.6),
        [
            "excess inventory put back to work",
            "",
            [(f"{F['stockout_units_avoided']:,.0f} lost-sale units avoided", {"bold": True})],
        ],
        size=16,
        color=BLACK,
    )
    rows = [
        ["Measure", "Before", "After"],
        ["Excess above target", money(F["excess_before"]), money(F["excess_after"])],
        ["Annual carrying cost of excess", money(F["carry_before"], 2), money(F["carry_after"], 2)],
        [
            "Expected lost-sale units (lead time)",
            f"{F['stockout_units_before']:,.0f}",
            f"{F['stockout_units_after']:,.0f}",
        ],
        ["Positions at stock-out risk", f"{F['risk_positions']:,}", f"{F['risk_after']:,}"],
        ["Inventory turns", f"{F['turns_before']:.1f}", f"{F['turns_after']:.1f}"],
        ["Net savings (one cycle)", "–", money(F["net_savings"])],
    ]
    R, L = PP_ALIGN.RIGHT, PP_ALIGN.LEFT
    table(
        s,
        Inches(5.15),
        BODY_TOP + Inches(0.1),
        Inches(7.58),
        rows,
        [Inches(4.18), Inches(1.7), Inches(1.7)],
        size=13,
        row_h=Inches(0.56),
        align=[L, R, R],
    )
    text(
        s,
        Inches(5.15),
        Inches(6.2),
        Inches(7.58),
        Inches(0.5),
        "Model estimates for one planning cycle versus doing nothing.",
        size=11,
        color=GRAY,
    )


def s7_dashboard(prs):
    """Slide 7: dashboard screenshots with callouts."""
    s = frame(
        prs,
        "A branch manager sees what to move, why, and approves it on one screen",
        7,
        f"""
This is the dashboard planners would actually use. It's built for branch and inventory teams, not data scientists: business language, every number in dollars, units or days.

Callout one: the page opens with the decision, in a sentence: how many units to move, what it costs and what it saves.

Callout two: the corridor map shows where stock is flowing, so a regional manager can sanity-check the plan at a glance.

Callout three, on the right: every recommended transfer is a row. The planner sets Approve or Reject and can add a note. That decision is recorded in an audit log with their name, the time and which model run it came from. Each row also carries a plain-language reason, for example: "{F['top_transfers'][0]['reason_text'][:150]}..."

Other pages show inventory health by branch and category, a forecast explorer for any branch and product with the drivers explained in plain language, and an accuracy page that reports bias by branch so nobody is quietly short-changed.

There's a live, view-only demo at {DEMO_URL}.

Likely question: "Can it execute transfers in the ERP?" Answer: deliberately not. It recommends; people decide and then raise the transfer. That keeps accountability with the planner.
""",
    )
    gutter = Inches(0.45)  # room for marker 1 left of the screenshot
    a = picture(s, A / "dash_top.png", ML + gutter, BODY_TOP, Inches(5.9), Inches(4.75))
    b = picture(s, A / "dash_table.png", Inches(7.3), BODY_TOP, Inches(5.43), Inches(3.3))
    # marker positions as fractions of each screenshot
    callout(s, 1, a.left - Inches(0.27), a.top + int(a.height * 0.10))
    callout(s, 2, a.left + int(a.width * 0.47), a.top + int(a.height * 0.69))
    callout(s, 3, b.left + int(b.width * 0.39), b.top - Inches(0.05))
    labels = [("1", "The plan in one sentence"), ("2", "Where stock flows"), ("3", "Approve or reject; logged")]
    y = b.top + b.height + Inches(0.25)
    for i, (n, lab) in enumerate(labels):
        callout(s, n, b.left + Inches(0.18), y + Inches(0.18) + i * Inches(0.42), d=Inches(0.34))
        text(
            s,
            b.left + Inches(0.5),
            y + Inches(0.03) + i * Inches(0.42),
            Inches(4.8),
            Inches(0.34),
            lab,
            size=13,
            color=BLACK,
        )
    text(
        s,
        ML + gutter,
        a.top + a.height + Inches(0.08),
        Inches(5.9),
        Inches(0.3),
        f"Live demo (view-only): {DEMO_URL}",
        size=11,
        color=GRAY,
    )


def s8_next(prs):
    """Slide 8: pilot plan, limitations and the mission."""
    s = frame(
        prs,
        "Next: a 12-week pilot at 2–3 branches on real ERP and warehouse data",
        8,
        """
To move from a demonstration to a decision, I'd propose a 12-week pilot.

Weeks one to four: connect real data. The tool reads seven standard extracts (sales, stock movements, open orders, promotions, branch and product masters, lane costs), all documented in the repository with validation rules, so a bad file is rejected with a clear message instead of a wrong answer.

Weeks five to ten: run it in shadow mode at two or three branches. Planners see the recommendations, approve or reject them, and we compare what the tool suggested with what actually happened. Rejections are valuable: they tell us where the cost assumptions are wrong.

Weeks eleven and twelve: decide. Monitoring is already built: weekly checks of forecast error and bias against the baseline, data drift, and an alert that triggers retraining.

The limitations, stated plainly: the data so far is synthetic; the plan looks one replenishment cycle ahead, not a multi-week schedule; safety stock assumes a normal error distribution; and the cost settings, like the price of a lost sale, need sign-off from finance before any dollar figure is trusted.

To close where we started: getting the right product to the right branch is one concrete way we make healthcare better with our customers. Surgeons don't wait, so patients don't either. Thank you; I'm happy to take questions.

Likely question: "What would you need from us?" Answer: read-only extracts for two or three branches, a planner willing to review recommendations weekly, and thirty minutes with finance on the cost settings.
""",
    )
    phases = [
        ("Weeks 1–4", "Connect data", "7 standard extracts,\nvalidated on load"),
        ("Weeks 5–10", "Shadow pilot", "2–3 branches; planners\napprove or reject"),
        ("Weeks 11–12", "Decide", "Monitor error and bias;\nretrain on drift"),
    ]
    n = len(phases)
    gap = Inches(0.45)
    bw = (CW - gap * (n - 1)) // n
    y = BODY_TOP + Inches(0.1)
    for i, (when, what, sub) in enumerate(phases):
        x = ML + i * (bw + gap)
        rect(s, x, y, bw, Inches(0.09), GOLD if i == 1 else CHARCOAL)
        text(s, x, y + Inches(0.22), bw, Inches(0.35), when, size=12, color=GRAY, bold=True)
        text(s, x, y + Inches(0.55), bw, Inches(0.5), what, size=20, color=BLACK, bold=True)
        text(s, x, y + Inches(1.05), bw, Inches(0.8), sub.split("\n"), size=13, color=CHARCOAL)
        if i < n - 1:
            arrow(s, x + bw + Inches(0.08), y + Inches(0.045), x + bw + gap - Inches(0.08), y + Inches(0.045))
    rect(s, ML, Inches(4.25), CW, Emu(9525), LIGHT)
    text(s, ML, Inches(4.45), Inches(3), Inches(0.35), "Limitations", size=14, bold=True, color=BLACK)
    lims = [
        "Synthetic data so far",
        "Plans one replenishment cycle",
        "Normal error model for safety stock",
        "Costs need finance approval",
    ]
    cw = (CW - Inches(0.3) * 3) // 4
    for i, lim in enumerate(lims):
        x = ML + i * (cw + Inches(0.3))
        box(s, x, Inches(4.9), cw, Inches(0.75), lim, fill=LIGHTER, line=LIGHT, size=13, bold=False, color=BLACK)
    text(
        s,
        ML,
        Inches(6.0),
        CW,
        Inches(0.6),
        "Right product, right branch: surgeons don't wait, so patients don't either.",
        size=16,
        bold=True,
        color=BLACK,
        align=PP_ALIGN.CENTER,
    )


def main() -> None:
    """Build the deck from assets/ and facts.json."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for fn in (s1_title, s2_problem, s3_approach, s4_forecast, s5_rebalance, s6_impact, s7_dashboard, s8_next):
        fn(prs)
    out = HERE / "Stockflow_Stryker.pptx"
    prs.save(out)
    print(f"wrote {out.relative_to(ROOT)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
